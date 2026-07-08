from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "JEI Manuscript Template v.7.2025.docx"
TEX_PATH = BASE / "main.tex"
OUT = BASE / "MEN2_Predictor_JEI_submission.docx"

SHORT_TITLE = (
    "Low-cost machine-learning risk stratification for medullary thyroid "
    "carcinoma in MEN2/RET carriers"
)


def section_chunk(tex, name, starred=False):
    star = r"\*" if starred else r"\*?"
    pattern = (
        rf"\\section{star}\{{{re.escape(name)}\}}([\s\S]*?)"
        rf"(?=\\section\*?\{{|\\appendix|\\end\{{document\}})"
    )
    match = re.search(pattern, tex)
    return match.group(1).strip() if match else ""


def subsection_chunks(section_text):
    parts = re.split(r"\\subsection\{([^}]+)\}", section_text)
    leading = parts[0].strip()
    chunks = []
    for i in range(1, len(parts), 2):
        chunks.append((parts[i].strip(), parts[i + 1].strip()))
    return leading, chunks


def strip_tables(text):
    text = re.sub(r"\\begin\{table\}[\s\S]*?\\end\{table\}", "\n\n", text)
    text = re.sub(r"\\begin\{longtable\}[\s\S]*?\\end\{longtable\}", "\n\n", text)
    return text


def latex_clean(text):
    text = strip_tables(text)
    text = text.replace("\\%", "%").replace("\\&", "&").replace("\\_", "_")
    text = re.sub(r"(?<!\\)%.*", "", text)
    text = text.replace("\\clearpage", "").replace("\\appendix", "")
    text = text.replace("---", "-").replace("--", "-")
    text = re.sub(r"\\emph\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\textbf\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\url\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\doi\{([^{}]*)\}", r"https://doi.org/\1", text)
    text = re.sub(r"\\weblink\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\href\{mailto:([^{}]*)\}\{([^{}]*)\}", r"\2", text)
    text = re.sub(r"\\href\{([^{}]*)\}\{([^{}]*)\}", r"\2", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?(?:\{[^{}]*\})?", "", text)
    text = text.replace("~", " ")
    text = re.sub(r"\$\^\{?([a-zA-Z0-9,*]+)\}?\$", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def paragraphs_from(text):
    text = strip_tables(text)
    raw = re.split(r"\n\s*\n", text.strip())
    return [latex_clean(paragraph) for paragraph in raw if latex_clean(paragraph)]


SUMMARY = (
    "Multiple Endocrine Neoplasia type 2 (MEN2) is a rare inherited cancer "
    "syndrome in which RET genetic testing is central to care but may be "
    "difficult to access in resource-limited settings. We hypothesized that "
    "routine clinical and biomarker variables could support high-sensitivity "
    "risk stratification for medullary thyroid carcinoma (MTC) in published "
    "MEN2/RET-carrier records without using RET variant as a model input. We "
    "curated 149 RET carriers from 10 peer-reviewed sources, harmonized "
    "demographic, clinical, biomarker, and genotype variables, and evaluated "
    "logistic regression, random forest, XGBoost, LightGBM, and linear support "
    "vector machine models using an 80/20 stratified internal held-out split. "
    "The main genotype-blind XGBoost model removed RET variant and ATA risk "
    "features and retained 93.3% sensitivity with 73.3% accuracy. A "
    "sequencing-informed genotype-aware XGBoost comparator reached 100.0% "
    "sensitivity with 83.3% accuracy. A secondary synthetic-augmentation "
    "experiment reached 96.2% accuracy with LightGBM, but this result was "
    "interpreted only as a simulation because synthetic controls may not "
    "represent real clinical heterogeneity. These findings do not provide a "
    "deployable diagnostic tool. Instead, they provide an open rare-disease "
    "dataset and reproducible benchmark suggesting that routine clinical and "
    "biomarker features may help prioritize future prospective MEN2 triage "
    "studies while highlighting important limitations from small sample size, "
    "imputation, biomarker timing, and lack of external validation."
)


REFERENCES = [
    "Wells, Samuel A., et al. “Revised American Thyroid Association Guidelines for the Management of Medullary Thyroid Carcinoma.” Thyroid, vol. 25, no. 6, 2015, pp. 567-610. https://doi.org/10.1089/thy.2014.0335",
    "Frank-Raue, K., et al. “Long-term Outcome in 46 Gene Carriers of Hereditary Medullary Thyroid Carcinoma after Prophylactic Thyroidectomy: Impact of Individual RET Genotype.” European Journal of Endocrinology, vol. 155, no. 2, 2006, pp. 229-236. https://doi.org/10.1530/eje.1.02216",
    "Xu, Jian Yu, et al. “Medullary Thyroid Carcinoma Associated with Germline RET K666N Mutation.” Thyroid, vol. 26, no. 12, 2016, pp. 1744-1751. https://doi.org/10.1089/thy.2016.0374",
    "Schulte, K. M., et al. “The Clinical Spectrum of Multiple Endocrine Neoplasia Type 2a Caused by the Rare Intracellular RET Mutation S891A.” Journal of Clinical Endocrinology & Metabolism, vol. 95, no. 9, 2010, pp. E92-E97. https://doi.org/10.1210/jc.2010-0375",
    "Qi, Xiao-Ping, et al. “RET Mutation p.S891A in a Chinese Family with Familial Medullary Thyroid Carcinoma and Associated Cutaneous Amyloidosis Binding OSMR Variant p.G513D.” Oncotarget, vol. 6, no. 32, 2015, pp. 33993-34003. https://doi.org/10.18632/oncotarget.4992",
    "Qi, X. P., et al. “The Rare Intracellular RET Mutation p.S891A in a Chinese Han Family with Familial Medullary Thyroid Carcinoma.” Journal of Biosciences, vol. 39, no. 3, 2014, pp. 505-512. https://doi.org/10.1007/s12038-014-9428-x",
    "Vijayan, Roopa, et al. “A Rare RET Mutation in an Indian Pedigree with Familial Medullary Thyroid Carcinoma.” Indian Journal of Cancer, vol. 58, no. 1, 2021, p. 98. https://doi.org/10.4103/ijc.IJC_639_19",
    "Florescu, Alexandru-Florin, et al. “Endocrine Perspective of Cutaneous Lichen Amyloidosis: RET-C634 Pathogenic Variant in Multiple Endocrine Neoplasia Type 2.” Clinics and Practice, vol. 14, no. 6, 2024, pp. 2284-2299. https://doi.org/10.3390/clinpract14060179",
    "La Greca, A., et al. “MEN2 Phenotype in a Family with Germline Heterozygous Rare RET K666N Variant.” Endocrinology, Diabetes & Metabolism Case Reports, vol. 2024, no. 3, 2024. https://doi.org/10.1530/EDM-24-0009",
    "Zhang, Hui-Fen, et al. “C634Y Mutation in RET-Induced Multiple Endocrine Neoplasia Type 2A: A Case Report.” World Journal of Clinical Cases, vol. 12, no. 15, 2024, pp. 2627-2635. https://doi.org/10.12998/wjcc.v12.i15.2627",
    "Shankar, R. K., et al. “Medullary Thyroid Cancer in a 9-week-old Infant with Familial MEN 2B: Implications for Timing of Prophylactic Thyroidectomy.” International Journal of Pediatric Endocrinology, vol. 2012, 2012, p. 25. https://doi.org/10.1186/1687-9856-2012-25",
    "van Buuren, Stef, and Karin Groothuis-Oudshoorn. “mice: Multivariate Imputation by Chained Equations in R.” Journal of Statistical Software, vol. 45, no. 3, 2011, pp. 1-67. https://doi.org/10.18637/jss.v045.i03",
    "Chawla, N. V., et al. “SMOTE: Synthetic Minority Over-sampling Technique.” Journal of Artificial Intelligence Research, vol. 16, 2002, pp. 321-357. https://doi.org/10.1613/jair.953",
    "Pedregosa, Fabian, et al. “Scikit-learn: Machine Learning in Python.” Journal of Machine Learning Research, vol. 12, 2011, pp. 2825-2830. https://jmlr.org/papers/v12/pedregosa11a.html",
    "Chen, Tianqi, and Carlos Guestrin. “XGBoost: A Scalable Tree Boosting System.” Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, ACM, 2016, pp. 785-794. https://doi.org/10.1145/2939672.2939785",
    "Ke, Guolin, et al. “LightGBM: A Highly Efficient Gradient Boosting Decision Tree.” Advances in Neural Information Processing Systems, vol. 30, 2017. https://doi.org/10.48550/arXiv.1712.01043",
    "Lundberg, Scott M., and Su-In Lee. “A Unified Approach to Interpreting Model Predictions.” Advances in Neural Information Processing Systems, vol. 30, 2017. https://doi.org/10.48550/arXiv.1705.07874",
    "Licata, Lorena, et al. “A Rare Case of Negative Serum Calcitonin in Metastatic Medullary Thyroid Carcinoma: Diagnosis, Treatment, and Follow-Up Strategy.” American Journal of Case Reports, vol. 23, 2022, e935207. https://doi.org/10.12659/AJCR.935207",
    "DNA Labs India. “RET Gene RET, Selective Sequencing of Exons 5, 8, 10, 11 and 13-16 NGS Genetic DNA Test.” DNA Labs India. https://dnalabsindia.com/test/ret-gene-ret-selective-sequencing-of-exons-5-8-10-11-and-13-16-ngs-genetic-dna-test. Accessed 4 May 2026.",
    "HealthCheckup.co.in. “Calcitonin Test from Thyrocare.” HealthCheckup.co.in. https://www.healthcheckup.co.in/package/calcitonin. Accessed 5 May 2026.",
    "PharmEasy. “Carcinoembryonic Antigen (CEA) Test.” PharmEasy. https://pharmeasy.in/diagnostics/tests/cea-46. Accessed 5 May 2026.",
    "Zhang, X., et al. “The Machine Learning-based Model for Lateral Lymph Node Metastasis of Thyroid Medullary Carcinoma Improved the Prediction Ability of Occult Metastasis.” Cancer Medicine, vol. 13, 2024, e7155. https://doi.org/10.1002/cam4.7155",
    "Little, R. J. A. “Missing-data Adjustments in Large Surveys.” Journal of Business & Economic Statistics, vol. 6, no. 3, 1988, pp. 287-296. https://doi.org/10.1080/07350015.1988.10509663",
    "Collins, Gary S., et al. “TRIPOD+AI Statement: Updated Guidance for Reporting Clinical Prediction Models That Use Regression or Machine Learning Methods.” BMJ, vol. 385, 2024, e078378. https://doi.org/10.1136/bmj-2023-078378",
]


TABLES = [
    (
        "Table 1. Baseline characteristics of the literature-derived cohort stratified by MTC status.",
        ["Characteristic", "Overall (n=149)", "MTC present (n=73)", "MTC absent (n=76)"],
        [
            ["Age, mean years", "34.27", "43.30", "25.59"],
            ["Age, median years", "31.5", "43.0", "20.0"],
            ["Age range, years", "0.17-90.0", "0.17-85.0", "1.0-90.0"],
            ["Female sex, n (%)", "99 (66.4)", "46 (63.0)", "53 (69.7)"],
            ["Male sex, n (%)", "50 (33.6)", "27 (37.0)", "23 (30.3)"],
            ["Unique RET variants, n", "14", "11", "13"],
            ["Calcitonin elevated, n (%)", "58 (38.9)", "47 (64.4)", "11 (14.5)"],
            ["CEA imputed, n (%)", "137 (91.9)", "65 (89.0)", "72 (94.7)"],
            ["Thyroid nodules present, n (%)", "10 (6.7)", "8 (11.0)", "2 (2.6)"],
            ["Family history of MTC, n (%)", "105 (70.5)", "53 (72.6)", "52 (68.4)"],
            ["Pheochromocytoma, n (%)", "9 (6.0)", "9 (12.3)", "0 (0.0)"],
            ["Hyperparathyroidism, n (%)", "5 (3.4)", "4 (5.5)", "1 (1.3)"],
        ],
    ),
    (
        "Table 2. Published sources used to construct the literature-derived cohort.",
        ["Study", "Source", "Main variant context", "n", "DOI"],
        [
            ["1", "Xu et al., 2016 (3)", "RET K666N families", "24", "https://doi.org/10.1089/thy.2016.0374"],
            ["2", "Frank-Raue et al., 2006 (2)", "Multiple RET variants after thyroidectomy", "46", "https://doi.org/10.1530/eje.1.02216"],
            ["3", "Qi et al., 2015 (5)", "RET S891A/R525W familial MTC", "15", "https://doi.org/10.18632/oncotarget.4992"],
            ["4", "Florescu et al., 2024 (8)", "RET C634G kindred", "6", "https://doi.org/10.3390/clinpract14060179"],
            ["5", "La Greca et al., 2024 (9)", "RET K666N family", "4", "https://doi.org/10.1530/EDM-24-0009"],
            ["6", "Vijayan et al., 2021 (7)", "Indian RET S891A pedigree", "7", "https://doi.org/10.4103/ijc.IJC_639_19"],
            ["7", "Zhang et al., 2024 (10)", "RET C634Y family", "3", "https://doi.org/10.12998/wjcc.v12.i15.2627"],
            ["8", "Shankar et al., 2012 (11)", "Familial MEN2B infant", "2", "https://doi.org/10.1186/1687-9856-2012-25"],
            ["9", "Schulte et al., 2010 (4)", "RET S891A multicenter cohort", "36", "https://doi.org/10.1210/jc.2010-0375"],
            ["10", "Qi et al., 2014 (6)", "Chinese Han RET S891A family", "6", "https://doi.org/10.1007/s12038-014-9428-x"],
        ],
    ),
    (
        "Table 3. Key variable groups used in the benchmark.",
        ["Variable group", "Variables"],
        [
            ["Demographic", "Age, sex, age group, age squared"],
            ["Genotype-aware", "RET variant one-hot encodings, ATA-aligned risk level, risk-age and risk-calcitonin interactions"],
            ["Biomarker", "Calcitonin elevated flag, numeric calcitonin, CEA numeric value, CEA missingness/imputation flag"],
            ["Clinical presentation", "Thyroid nodule indicators, family history of MTC, pheochromocytoma, hyperparathyroidism"],
            ["Outcome", "Binary MTC status extracted from source reports"],
        ],
    ),
    (
        "Table 4. Selected missingness and harmonization summary in the 149-record literature-derived cohort.",
        ["Field", "Observed or nonzero", "Total", "Note"],
        [
            ["Age", "149", "149", "Missing ages were filled from source-level medians where required."],
            ["Sex", "149", "149", "Unknown values were mode-filled for modeling."],
            ["RET variant", "149", "149", "Required for genotype-aware inclusion."],
            ["Calcitonin numeric value", "77", "149", "Zero includes unavailable, normal, or undetectable values after harmonization."],
            ["CEA observed before imputation", "12", "149", "Observed paired CEA values seeded the imputation analysis."],
            ["CEA imputed", "137", "149", "Missingness flag retained in the model dataset."],
        ],
    ),
    (
        "Table 5. Internal held-out performance with raw counts and Wilson confidence intervals.",
        ["Analysis", "Model", "Confusion matrix", "Sensitivity", "Accuracy"],
        [
            ["Genotype-aware", "XGBoost", "TN=10, FP=5, FN=0, TP=15", "100.0% (15/15; 79.6%-100.0%)", "83.3% (25/30; 66.4%-92.7%)"],
            ["Genotype-blind", "XGBoost", "TN=8, FP=7, FN=1, TP=14", "93.3% (14/15; 70.2%-98.8%)", "73.3% (22/30; 55.6%-85.8%)"],
            ["No CEA", "XGBoost", "TN=12, FP=3, FN=0, TP=15", "100.0% (15/15; 79.6%-100.0%)", "90.0% (27/30; 74.4%-96.5%)"],
            ["No biomarkers", "XGBoost", "TN=12, FP=3, FN=0, TP=15", "100.0% (15/15; 79.6%-100.0%)", "90.0% (27/30; 74.4%-96.5%)"],
            ["Synthetic", "LightGBM", "TN=156, FP=3, FN=5, TP=46", "90.2% (46/51; 79.0%-95.7%)", "96.2% (202/210; 92.7%-98.1%)"],
        ],
    ),
]


FIGURES = [
    (
        BASE / "figures_jpeg" / "Figure_1_variant_distribution.jpg",
        "Figure 1. RET variant distribution in the literature-derived MEN2/RET-carrier cohort. The chart summarizes the number of curated carrier records for each RET variant represented in the 149-record dataset.",
    ),
    (
        BASE / "figures_jpeg" / "Figure_2_age_distribution.jpg",
        "Figure 2. Age distribution of RET carriers in the literature-derived cohort. The histogram shows the age spread of the curated records used for internal model benchmarking.",
    ),
    (
        BASE / "figures_jpeg" / "Figure_3_calcitonin_cea_relationship.jpg",
        "Figure 3. Observed calcitonin-CEA relationship in records with paired biomarker values. Only 12 paired observations were available, so CEA analyses were treated as exploratory.",
    ),
]


def configure_paragraph(paragraph, indent=True, align=None):
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5) if indent else None
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(11)


def add_paragraph(doc, text="", indent=True, align=None, bold=False, italic=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(11)
    run.bold = bold
    run.italic = italic
    configure_paragraph(paragraph, indent=indent, align=align)
    return paragraph


def add_heading(doc, text):
    return add_paragraph(doc, text.upper(), indent=False, bold=True)


def add_subheading(doc, text):
    return add_paragraph(doc, text, indent=False, italic=True)


def add_page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    configure_paragraph(paragraph, indent=False)


def add_caption(doc, caption):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = None
    match = re.match(r"^(Figure|Table) (\d+\.)\s*(.*)$", caption)
    if match:
        first = paragraph.add_run(f"{match.group(1)} {match.group(2)} ")
        first.bold = True
        second = paragraph.add_run(match.group(3))
        runs = (first, second)
    else:
        first = paragraph.add_run(caption)
        first.bold = True
        runs = (first,)
    for run in runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(11)


def set_cell_text(cell, text, bold=False, size=9):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=9)
    return table


def build():
    tex = TEX_PATH.read_text(encoding="utf-8")

    intro_paragraphs = paragraphs_from(section_chunk(tex, "Introduction"))
    intro_paragraphs = intro_paragraphs[:3] + [
        "We hypothesized that routine clinical and biomarker variables would allow high-sensitivity MTC risk stratification in published MEN2/RET-carrier records without using RET variant as an input. To test this hypothesis, we separated two modeling tracks. The genotype-blind analysis removed RET variant and American Thyroid Association (ATA) risk features and was treated as the clinically motivated triage benchmark because it relied on non-genetic clinical and biomarker variables. The genotype-aware benchmark included RET variant and ATA-aligned risk features and was interpreted as a sequencing-informed comparator relevant only after genetic context was already known. A secondary synthetic-augmentation experiment was reported separately as a simulation, not as an expanded real-world cohort."
    ]

    methods_leading, methods_chunks = subsection_chunks(section_chunk(tex, "Materials and Methods"))
    results_leading, results_chunks = subsection_chunks(section_chunk(tex, "Results"))
    discussion_text = section_chunk(tex, "Discussion")
    limitations_text = ""
    if r"\subsection{Limitations}" in discussion_text:
        discussion_main, limitations_text = discussion_text.split(r"\subsection{Limitations}", 1)
    else:
        discussion_main = discussion_text
    discussion_paragraphs = (
        paragraphs_from(discussion_main)
        + paragraphs_from(limitations_text)
        + paragraphs_from(section_chunk(tex, "Conclusion"))
    )
    acknowledgments = paragraphs_from(section_chunk(tex, "Acknowledgements", starred=True))
    acknowledgments.append(
        "This research received no specific funding from any funding agency in the public, commercial, or not-for-profit sectors."
    )
    data_availability = paragraphs_from(section_chunk(tex, "Data Availability", starred=True))[0]
    methods_chunks.append(("Data and code availability", data_availability))

    doc = Document(TEMPLATE)
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    doc.styles["Normal"].font.size = Pt(11)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(old)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    sect_pr.append(line_numbers)

    title = add_paragraph(doc, SHORT_TITLE, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    title.runs[0].font.size = Pt(12)
    add_paragraph(doc, "", indent=False)
    add_paragraph(
        doc,
        "Harnoor Kaur1, Arjun Vijay Prakash1, Shashwat Mishra1",
        indent=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(
        doc,
        "1 City Montessori School, Lucknow, Uttar Pradesh, India",
        indent=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(doc, "", indent=False)
    add_paragraph(doc, "Student Authors", indent=False, bold=True)
    add_paragraph(doc, "Harnoor Kaur, high school", indent=False)
    add_paragraph(doc, "Arjun Vijay Prakash, high school", indent=False)
    add_paragraph(doc, "", indent=False)
    add_paragraph(doc, "KEYWORDS: MEN2; RET; calcitonin; machine learning; screening", indent=False)
    add_paragraph(doc, "", indent=False)
    add_paragraph(
        doc,
        "OVERVIEW: This manuscript presents an open MEN2/RET-carrier dataset and a reproducible machine-learning benchmark for MTC risk stratification. The main genotype-blind model suggests that routine clinical and biomarker variables may support future high-sensitivity triage research, but prospective validation is required before clinical use.",
        indent=False,
    )
    add_page_break(doc)

    add_heading(doc, "Summary")
    add_paragraph(doc, SUMMARY)
    add_page_break(doc)

    add_heading(doc, "Introduction")
    for paragraph in intro_paragraphs:
        add_paragraph(doc, paragraph)

    add_heading(doc, "Results")
    if results_leading:
        for paragraph in paragraphs_from(results_leading):
            add_paragraph(doc, paragraph)
    for title, chunk in results_chunks:
        add_subheading(doc, title)
        for paragraph in paragraphs_from(chunk):
            add_paragraph(doc, paragraph)

    add_heading(doc, "Discussion")
    for paragraph in discussion_paragraphs:
        add_paragraph(doc, paragraph)

    add_heading(doc, "Materials and Methods")
    if methods_leading:
        for paragraph in paragraphs_from(methods_leading):
            add_paragraph(doc, paragraph)
    for title, chunk in methods_chunks:
        add_subheading(doc, title)
        for paragraph in paragraphs_from(chunk):
            add_paragraph(doc, paragraph)

    add_heading(doc, "Acknowledgments")
    for paragraph in acknowledgments:
        add_paragraph(doc, paragraph)

    add_heading(doc, "References")
    for index, reference in enumerate(REFERENCES, 1):
        add_paragraph(doc, f"{index}. {reference}", indent=False)

    add_heading(doc, "Figures and Figure Titles/Captions")
    for path, caption in FIGURES:
        doc.add_picture(str(path), width=Inches(6.25))
        image_paragraph = doc.paragraphs[-1]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        configure_paragraph(image_paragraph, indent=False)
        add_caption(doc, caption)

    add_heading(doc, "Tables with Titles/Captions")
    for caption, headers, rows in TABLES:
        add_table(doc, headers, rows)
        add_caption(doc, caption)

    doc.core_properties.title = SHORT_TITLE
    doc.core_properties.subject = "JEI manuscript submission"
    doc.core_properties.author = "Harnoor Kaur; Arjun Vijay Prakash; Shashwat Mishra"
    doc.save(OUT)

    with zipfile.ZipFile(OUT) as archive:
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = required.difference(archive.namelist())
        if missing:
            raise RuntimeError(f"Missing required DOCX parts: {missing}")
    Document(OUT)

    summary_words = len(re.findall(r"\b[\w%-]+\b", SUMMARY))
    print(OUT)
    print(f"Title chars: {len(SHORT_TITLE)}")
    print(f"Summary words: {summary_words}")
    print(f"Figures: {len(FIGURES)} Tables: {len(TABLES)} Total: {len(FIGURES) + len(TABLES)}")


if __name__ == "__main__":
    build()
